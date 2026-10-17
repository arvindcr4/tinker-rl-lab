#!/usr/bin/env python3
"""Generate an evidence-linked status deck; never refresh historical fixed claims.

The archived deck is not overwritten. Pass a validated structured result file
explicitly. Reading old progress logs establishes only the last logged state,
not a currently running process. Missing logs remain unknown.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import hashlib
from pathlib import Path
import re
import sys

sys.path.append(str(Path(__file__).resolve().parent))
from _paths import REPO_ROOT
from _review02_contracts import result_provenance

ROOT = REPO_ROOT


def _read(path):
    try:
        return Path(path).read_text(errors='replace'), None
    except OSError as exc:
        return '', str(exc)


def log_status(path, pattern):
    text, error = _read(path)
    if error:
        return f'unknown: {error}'
    progress = {}
    for arm, step, total in re.findall(pattern, text):
        progress[arm] = (int(step), int(total))
    if not progress:
        return 'unknown: no parseable progress records'
    return 'last logged, current liveness unknown: ' + '; '.join(f'{arm} {step}/{total}' for arm, (step, total) in sorted(progress.items()))


def n2_status():
    return log_status(ROOT/'platform_hybrid/experiments/tinker-runs/logs/n2_reward_tensor_resume_20260704.out',
                      r'\[n2:(\w+)\]\s+(\d+)/(\d+)')


def n10_status():
    return log_status(ROOT/'platform_hybrid/experiments/tinker-runs/logs/n10_gsm8k_cot_seed_expansion_20260704.out',
                      r'\[(n10_\w+?_s\d+)\]\s+(\d+)/(\d+)')


def mega_status():
    text, error = _read(ROOT/'platform_hybrid/experiments/results/mega_20260704/cells_done.jsonl')
    if error:
        return f'unknown: {error}'
    try:
        records = [json.loads(line) for line in text.splitlines() if line.strip()]
    except ValueError as exc:
        return f'unknown: invalid completion log: {exc}'
    return f'{len(records)} logged completion records; current liveness unknown'


def validated_results(path):
    data = json.loads(Path(path).read_text())
    rows = data.get('results')
    required = {'estimand', 'estimate', 'ci_low', 'ci_high', 'unit', 'method', 'status',
                'input_sha256', 'analysis_commit', 'statistic_version', 'sample_counts', 'exclusions',
                'seed', 'uncertainty_source'}
    if not isinstance(rows, list) or not rows:
        raise ValueError('nonempty structured results required')
    for row in rows:
        if not isinstance(row, dict) or not required.issubset(row):
            raise ValueError('every result needs estimand, interval, status, counts, exclusions and reproducibility metadata')
        if not isinstance(row['input_sha256'], dict) or not row['input_sha256'] or not re.fullmatch(r'[0-9a-f]{40}', row['analysis_commit']):
            raise ValueError('verified input and analysis identities required')
        for source, digest in row['input_sha256'].items():
            source_path = Path(source)
            if not source_path.is_absolute():
                source_path = ROOT/source_path
            if not source_path.is_file() or hashlib.sha256(source_path.read_bytes()).hexdigest() != digest:
                raise ValueError(f'input hash cannot be verified: {source}')
    return rows


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--results', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args(argv)
    results = validated_results(args.results)
    from pptx import Presentation
    prs = Presentation()

    def slide(title, text):
        page = prs.slides.add_slide(prs.slide_layouts[1])
        page.shapes.title.text = title
        page.placeholders[1].text = text

    slide('ZVF: scope of the measurement', 'ZVF measures absence of within-group reward contrast. It does not establish collapse, absence of KL gradients, or absence of learning from other groups.\nGenerated '+datetime.now(timezone.utc).isoformat())
    slide('Last logged experiment state', '\n'.join([n2_status(), n10_status(), mega_status()]))
    for row in results:
        slide(row['estimand'], f"Estimate: {row['estimate']} {row['unit']}\nInterval: [{row['ci_low']}, {row['ci_high']}]\nStatus: {row['status']}\nMethod: {row['method']}\nUncertainty: {row['uncertainty_source']}\nAnalysis commit: {row['analysis_commit']}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    manifest = result_provenance([args.results], results=results, generation_time=datetime.now(timezone.utc).isoformat())
    Path(str(args.output)+'.manifest.json').write_text(json.dumps(manifest, indent=2, allow_nan=False))


if __name__ == '__main__':
    main()
